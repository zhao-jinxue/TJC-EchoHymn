#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""软著源代码清洗脚本（EchoHymn ruanzhu-workspace / 01 提示词 Step 2 前置工具）。

模式：
  report  —— Step 1 资产盘点：逐文件统计原始/清洗后行数 + 敏感信息预扫描（不改写任何文件）
  build   —— Step 2 产物生成：按传入的有序文件清单清洗拼接，取前 N 页 + 后 M 页写入输出文件，
              并做每页 50 行与总行数断言校验

清洗规则（对齐 ruanzhu-workspace/01_cline_master_prompt.md）：
  删除空行与纯注释行；删除行尾 // 注释与 /*..*/ 块注释（支持 Dart 嵌套块注释）；
  字符串字面量（含 raw / 三引号多行串）内的 // 与引号安全处理；
  import/library/part/export 指令行原样保留；行尾空白去除；统一 \n。
"""
import io
import json
import re
import sys
from pathlib import Path

PAGE_LINES = 50


def strip_dart(src: str) -> list[str]:
    """返回清洗后的行列表。逐字符状态机：注释/空行剔除，字符串内容原样保留。"""
    out: list[str] = []
    buf: list[str] = []
    block_depth = 0           # /* */ 嵌套深度（Dart 支持嵌套块注释）
    quote: str | None = None  # 当前字符串定界符: ' / " / ''' / """
    raw = False               # 当前字符串是否 r 前缀
    n = len(src)
    i = 0
    while i < n:
        ch = src[i]
        if block_depth:
            if src.startswith("/*", i):
                block_depth += 1
                i += 2
                continue
            if src.startswith("*/", i):
                block_depth -= 1
                i += 2
                continue
            if ch == "\n":
                if "".join(buf).strip():
                    out.append("".join(buf).rstrip())
                buf = []
            i += 1
            continue
        if quote:
            if not raw and ch == "\\" and quote in ("'", '"') and len(quote) == 1:
                buf.append(src[i:i + 2])
                i += 2
                continue
            if src.startswith(quote, i):
                buf.append(quote)
                i += len(quote)
                quote = None
                continue
            if ch == "\n":  # 三引号串内换行：逐物理行 flush，保持"元素=物理行"不变量
                line = "".join(buf).rstrip()
                if line:
                    out.append(line)
                buf = []
                if len(quote) == 1:  # 单引号串内裸换行属异常，容错结束字符串
                    quote = None
                i += 1
                continue
            buf.append(ch)
            i += 1
            continue
        # 非字符串、非块注释：识别行注释 / 块注释 / 字符串起始 / 普通字符
        if src.startswith("//", i):  # 行注释（/// 文档注释同理被吞），跳到行尾
            j = src.find("\n", i)
            j = n if j < 0 else j
            i = j  # 换行符留给下一轮处理（触发 flush）
            continue
        if src.startswith("/*", i):
            block_depth = 1
            i += 2
            continue
        if ch in "rR" and i + 1 < n and src[i + 1] in "'\"":
            prev = src[i - 1] if i else " "
            if not (prev.isalnum() or prev in "_$"):  # raw 字符串前缀
                q3 = src[i + 1:i + 4]
                if q3 in ("'''", '"""'):
                    buf.append(src[i:i + 4])
                    i += 4
                    quote = q3
                    raw = True
                    continue
                buf.append(src[i:i + 2])
                i += 2
                quote = src[i - 1]
                raw = True
                continue
        if ch in "'\"":
            q3 = src[i:i + 3]
            if q3 in ("'''", '"""'):
                buf.append(q3)
                i += 3
                quote = q3
            else:
                buf.append(ch)
                i += 1
                quote = ch
            raw = False
            continue
        if ch == "\n":
            line = "".join(buf).rstrip()
            if line:
                out.append(line)
            buf = []
            i += 1
            continue
        buf.append(ch)
        i += 1
    tail = "".join(buf).rstrip()
    if tail:
        out.append(tail)
    return out


SENSITIVE = [
    ("IPv4", re.compile(r"\b(?!127\.0\.0\.1|0\.0\.0\.0|8\.8\.8\.8)(\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3})\b")),
    ("手机号", re.compile(r"\b1[3-9]\d{9}\b")),
    ("邮箱", re.compile(r"\b[\w.+-]+@[\w-]+\.[\w.]{2,}\b")),
    ("凭据字样", re.compile(r"(?i)\b(api[_-]?key|secret|token|passwd|password)\b\s*[:=]")),
    ("sk-密钥", re.compile(r"\bsk-[A-Za-z0-9]{8,}")),
    ("http链接", re.compile(r"https?://[^\s'\"\),]+")),
]

EXCLUDE = {"lib/data/chinese_convert_map.dart"}  # 01 规则 6：纯数据映射表


def collect(root: Path) -> list[Path]:
    files = sorted(p for p in root.rglob("*.dart"))
    return [p for p in files if str(p.relative_to(root.parent)).replace("\\", "/") not in EXCLUDE]


def rel(p: Path, root: Path) -> str:
    return str(p.relative_to(root.parent)).replace("\\", "/")


# ---- build 模式常量（对齐 output/material_plan.md §4，基线 commit 73b0d20）----
WS = Path(__file__).resolve().parent.parent            # ruanzhu-workspace/
APP = WS.parent / "hymn_app"
HEADER_TEXT = "EchoHymn 赞美诗播放软件 V1.5"
TOTAL_CLEANED = 5788           # rev1：三引号 SQL 拆物理行后 (sqlite_repository 238->246)
FRONT_LINES = 1500             # 前30页=换行后物理流行 1~1500；后30页=末 1500
UNIT_CAP = 53.0                # 9pt 版心 481.9pt -> 53.5 单位；保守 53.0 保证 Word 不再二次换行


def display_pt(line: str, font_pt: float) -> float:
    """行宽（磅）：Consolas 半宽 ≈0.55em，CJK 全宽 =1em。"""
    import unicodedata
    w = 0.0
    for ch in line.expandtabs(4):
        w += font_pt if unicodedata.east_asian_width(ch) in ("F", "W") else font_pt * 0.55
    return w


def wrap_stream(line: str) -> list[str]:
    """超宽行按保守宽度显式断行（终端换行式，不增删字符，tab 展开为 4 空格），
    使每个产物行在 Word 9pt 下一行一页内呈现——页数与"每页 50 行"由此完全确定。"""
    import unicodedata
    parts: list[str] = []
    cur: list[str] = []
    u = 0.0
    for ch in line.expandtabs(4):
        w = 1.0 if unicodedata.east_asian_width(ch) in ("F", "W") else 0.55
        if u + w > UNIT_CAP:
            parts.append("".join(cur))
            cur, u = [ch], w
        else:
            cur.append(ch)
            u += w
    if cur:
        parts.append("".join(cur))
    return parts


def make_docx(lines: list[str], path: Path, font_pt: float) -> None:
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2)
    s.header_distance = s.footer_distance = Cm(0.8)
    st = doc.styles["Normal"]
    st.font.name = "Consolas"
    st.font.size = Pt(font_pt)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = st.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    # 版心高 25.7cm=728.5pt ÷ 50 行 = 14.57pt -> 固定行距 14.5pt 精确保证每页 50 行
    pf.line_spacing = Pt(14.5)
    pf.widow_control = False
    hp = s.header.paragraphs[0]
    hp.text = HEADER_TEXT
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    for ln in lines:
        doc.add_paragraph(ln)
    doc.save(str(path))


def build() -> int:
    """Step 2：按 step2_file_plan 有序清单清洗拼接 -> 断言 -> 前1500+后1500 -> txt+docx 双产物。"""
    state = json.loads((WS / "progress_state.json").read_text(encoding="utf-8"))
    plan = state["step2_file_plan"]
    cleaned_total = 0
    wrapped: list[str] = []
    offsets: list[tuple[str, int, int]] = []      # (file, 换行流起, 换行流止) 1-based 由计算还原
    stream_head: list[str] = []                   # 仅用于逐行质检
    for e in plan:
        f = APP / Path(e["file"].replace("/", "\\"))
        cleaned = strip_dart(f.read_text(encoding="utf-8-sig", newline=""))
        if len(cleaned) != e["cleaned_lines"]:
            print(f"[FAIL] 漂移 {e['file']}: 计划={e['cleaned_lines']} 现状={len(cleaned)}"
                  f"（代码与已确认计划不一致，须重跑 Step 1）", file=sys.stderr)
            return 1
        start = len(wrapped)
        stream_head.extend(cleaned)
        for ln in cleaned:
            parts = wrap_stream(ln)
            wrapped.extend(parts)
            for p in parts:
                if display_pt(p, 1.0) > 53.5:
                    print(f"[FAIL] 断行后仍超宽: {p[:50]}", file=sys.stderr)
                    return 1
        offsets.append((e["file"], start + 1, len(wrapped)))
        cleaned_total += len(cleaned)
    if cleaned_total != TOTAL_CLEANED:
        print(f"[FAIL] 清洗后物理行总数 {cleaned_total} != 计划 {TOTAL_CLEANED}", file=sys.stderr)
        return 1
    for i, l in enumerate(stream_head, 1):
        if not l.strip():
            print(f"[FAIL] 空行残留 @{i}", file=sys.stderr); return 1
        if l.lstrip().startswith(("//", "/*")):
            print(f"[FAIL] 注释残留 @{i}: {l[:50]}", file=sys.stderr); return 1
    sens = [(i, pat.search(l).group(0)) for i, l in enumerate(stream_head, 1)
            for name, pat in SENSITIVE if pat.search(l)]
    if sens:
        print(f"[FAIL] 敏感信息 {len(sens)} 处: {sens[:5]}", file=sys.stderr); return 1
    out = wrapped[:FRONT_LINES] + wrapped[-FRONT_LINES:]
    if len(out) != 3000:
        print(f"[FAIL] 截取后 {len(out)} != 3000", file=sys.stderr); return 1
    imports = sum(1 for l in out if l.startswith(("import ", "export ", "part ", "library ")))
    if imports == 0:
        print("[FAIL] import 行数为 0", file=sys.stderr); return 1

    def seg_of(pos: int) -> str:
        for fname, a, b in offsets:
            if a <= pos <= b:
                return f"{fname}(换行流{a}~{b})"
        return "?"

    tail_start = len(wrapped) - FRONT_LINES + 1
    print(f"[INFO] 清洗 {cleaned_total} 行 -> 换行流 {len(wrapped)} 物理行（断行增 "
          f"{len(wrapped) - cleaned_total} 行）")
    print(f"[INFO] 前30页 = 换行流 1~1500，止于 {seg_of(FRONT_LINES)}")
    print(f"[INFO] 后30页 = 换行流 {tail_start}~{len(wrapped)}，起于 {seg_of(tail_start)}")
    (WS / "output" / "source_code.txt").write_text("\n".join(out) + "\n", encoding="utf-8", newline="\n")
    make_docx(out, WS / "output" / "source_code.docx", 9.0)
    from docx import Document as _D
    n = len(_D(str(WS / "output" / "source_code.docx")).paragraphs)
    if n != 3000:
        print(f"[FAIL] docx 段落数 {n} != 3000", file=sys.stderr); return 1
    print(f"[OK] source_code.txt = 3000 物理行（前 1500 + 后 1500）")
    print(f"[OK] source_code.docx：A4 / 2cm 边距 / Consolas 9pt+宋体东亚 / 固定行距 14.5pt"
          f"（版心 728.5pt÷50=14.57 -> 每页恰 50 行）/ 页眉「{HEADER_TEXT}」/ 页脚居中 PAGE 域 / "
          f"段落 {n} / import 行 {imports} / 敏感 {len(sens)} 命中")
    return 0


def main() -> int:
    mode = sys.argv[1] if len(sys.argv) > 1 else "report"
    root = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("E:/EchoHymn/hymn_app/lib")
    files = collect(root)
    total_raw = total_clean = 0
    hits: list[str] = []
    stats: dict[str, int] = {}

    for f in files:
        src = f.read_text(encoding="utf-8-sig", newline="")
        raw_lines = src.count("\n") + (0 if src.endswith("\n") else 1)
        cleaned = strip_dart(src)
        stats[rel(f, root)] = len(cleaned)
        total_raw += raw_lines
        total_clean += len(cleaned)
        for name, pat in SENSITIVE:
            for no, line in enumerate(cleaned, 1):
                m = pat.search(line)
                if m:
                    hits.append(f"{rel(f, root)}:{no} [{name}] {m.group(0)[:60]} | {line.strip()[:80]}")

    if mode == "report":
        print(f"{'file':<58}{'raw->?':>8}{'cleaned':>9}")
        for k in stats:
            print(f"{k:<58}{'':>8}{stats[k]:>9}")
        print(f"TOTAL include-scope files={len(files)} cleaned_lines={total_clean} "
              f"pages_needed={total_clean // PAGE_LINES} strategy={'60页(前30+后30)' if total_clean > 3000 else '全量'}")
        print(f"SENSITIVE_HITS={len(hits)}")
        for h in hits:
            print("  " + h)
        (root.parent / "../ruanzhu-workspace/tools_clean/stats.json").resolve().write_text(
            json.dumps({"total_cleaned": total_clean, "files": stats}, ensure_ascii=False, indent=2),
            encoding="utf-8")
        return 0

    if mode == "build":
        return build()

    print("usage: clean.py report [lib_dir] | build", file=sys.stderr)
    return 2


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    sys.exit(main())
