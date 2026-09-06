#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""软著用户手册 md -> docx 导出器（Step 3 配套，01 提示词 Step 3 产物排版）。

输入:  ../output/user_manual_draft.md（受控 Markdown 子集）
输出:  ../output/user_manual.docx —— A4 / 封面 / 目录域 / 黑体标题+宋体正文 /
       页眉全称+版本 / 页脚 PAGE 域 / 截图占位符红色居中醒目呈现。
"""
import io
import re
import sys
from pathlib import Path

WS = Path(__file__).resolve().parent.parent
SRC = WS / "output" / "user_manual_draft.md"
DST = WS / "output" / "user_manual.docx"
HEADER_TEXT = "EchoHymn 赞美诗播放软件 V1.5"
APPLICANT = "赵金雪"
DOC_DATE = "2026年9月"


def add_runs(p, text):
    """处理 **加粗** 与 `代码` 两类行内标记。"""
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|`([^`]+)`", text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        if m.group(1) is not None:
            r = p.add_run(m.group(1))
            r.bold = True
        else:
            r = p.add_run(m.group(2))
            r.font.name = "Consolas"
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])


def setup(doc):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.54)
    s.left_margin = s.right_margin = Cm(2.5)
    s.header_distance = s.footer_distance = Cm(1)
    st = doc.styles["Normal"]
    st.font.name = "宋体"
    st.font.size = Pt(12)  # 小四
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    st.paragraph_format.line_spacing = 1.5
    for name, size in (("Heading 1", 16), ("Heading 2", 14)):
        h = doc.styles[name]
        h.font.name = "黑体"
        h.font.size = Pt(size)
        h.font.bold = True
        rf = h.element.rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            h.element.rPr.append(rf)
        rf.set(qn("w:eastAsia"), "黑体")
    hp = s.header.paragraphs[0]
    hp.text = HEADER_TEXT
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)


def cover(doc):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    for _ in range(6):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("EchoHymn 赞美诗播放软件")
    r.bold = True
    r.font.size = Pt(26)
    v = doc.add_paragraph()
    v.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = v.add_run("V1.5")
    r.bold = True
    r.font.size = Pt(20)
    m = doc.add_paragraph()
    m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = m.add_run("用户操作手册")
    r.bold = True
    r.font.size = Pt(22)
    for _ in range(8):
        doc.add_paragraph()
    for line in (f"著作权人 / 申请人：{APPLICANT}", DOC_DATE):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line).font.size = Pt(14)
    doc.add_page_break()


def toc(doc):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    h = doc.add_paragraph()
    r = h.add_run("目  录")
    r.bold = True
    r.font.size = Pt(16)
    p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
    p._p.append(fld)
    tip = doc.add_paragraph()
    tip.add_run("（在 Word 中：Ctrl+A 后按 F9 更新目录域）").font.size = Pt(9)
    doc.add_page_break()
def flush_table(doc, rows):
    from docx.shared import Pt
    if not rows:
        return
    body = [r for r in rows if not re.match(r"^\|\s*-+", r)]
    ncol = max(len(r.strip("|").split("|")) for r in body)
    tbl = doc.add_table(rows=len(body), cols=ncol)
    tbl.style = "Table Grid"
    for i, raw in enumerate(body):
        cells = [c.strip() for c in raw.strip().strip("|").split("|")]
        for j in range(ncol):
            txt = cells[j] if j < len(cells) else ""
            cell = tbl.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, re.sub(r"\*\*(.+?)\*\*", r"\1", txt) if i == 0 else txt)
            if i == 0:
                for r in p.runs:
                    r.bold = True
            for r in p.runs:
                r.font.size = Pt(10.5)
    doc.add_paragraph()


def main() -> int:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    setup(doc)
    cover(doc)
    toc(doc)
    table_buf: list[str] = []
    n_ph = n_h = 0
    for ln in lines:
        raw = ln.rstrip()
        if raw.startswith("|"):
            table_buf.append(raw)
            continue
        flush_table(doc, table_buf)
        table_buf = []
        if not raw.strip() or raw == "---" or raw.startswith("# "):
            continue
        if raw.startswith("## "):
            doc.add_paragraph(raw[3:], style="Heading 1"); n_h += 1
        elif raw.startswith("### "):
            doc.add_paragraph(raw[4:], style="Heading 2"); n_h += 1
        elif raw.startswith("[此处插入") and raw.endswith("]"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"〔 截图占位 {raw} 〕")
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            r.bold = True
            n_ph += 1
        elif re.match(r"^\s*(-|\d+\.)\s", raw):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(24)
            add_runs(p, raw.strip())
        else:
            p = doc.add_paragraph()
            add_runs(p, raw.strip())
    flush_table(doc, table_buf)
    doc.save(str(DST))
    d = Document(str(DST))
    print(f"[OK] user_manual.docx 生成：段落 {len(d.paragraphs)}，标题 {n_h}，"
          f"截图占位 {n_ph}，表格 {len(d.tables)}")
    assert n_h >= 20 and n_ph >= 15 and len(d.tables) >= 3, "手册结构不完整"
    return 0


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    sys.exit(main())
