# -*- coding: utf-8 -*-
"""软著「软件说明书（操作手册）」生成器（Word，V 版本取自 pubspec 单源）。

体例：封面信息 → 概述 → 安装 → 界面总览 → 各功能操作 → 数据与文件 → 版本信息 → 版权声明；
页眉 = 软件名称 + 版本号 + 页码域（与源程序鉴别材料同口径）。
插图位于 docs/ruanzhu/v<版本>/img/（组件级版式图；实机整屏截图建议后续补充，
见同目录《材料清单与待补项》）。

用法：python tools/make_ruanzhu_manual.py
"""
import pathlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUTDIR = ROOT / 'docs' / 'ruanzhu'
APP_NAME = 'EchoHymn · 聆听赞美诗'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def app_version():
    for line in (ROOT / 'hymn_app' / 'pubspec.yaml').read_text(encoding='utf-8').splitlines():
        if line.startswith('version:'):
            return line.split(':', 1)[1].strip().split('+')[0].strip()
    raise SystemExit('无法解析 pubspec 版本')


def _fld_run(instr):
    r = OxmlElement('w:r')
    for tag, attr, text in (
            ('w:fldChar', 'begin', None), ('w:instrText', None, instr),
            ('w:fldChar', 'separate', None), ('w:fldChar', 'end', None)):
        el = OxmlElement(tag)
        if attr:
            el.set(W + 'fldCharType', attr)
        if text:
            el.set(W + 'space', 'preserve')
            el.text = text
        r.append(el)
    return r


def build(version):
    img = OUTDIR / f'v{version}' / 'img'
    doc = Document()
    sec = doc.sections[0]
    header = sec.header.paragraphs[0]
    header.text = f'{APP_NAME} 软件说明书  V{version}    '
    header.add_run('第 ')
    header._p.append(_fld_run(' PAGE '))
    header.add_run(' 页 / 共 ')
    header._p.append(_fld_run(' NUMPAGES '))
    header.add_run(' 页')
    for run in header.runs:
        run.font.size = Pt(9)
    normal = doc.styles['Normal']
    normal.font.name = '宋体'
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(W + 'eastAsia', '宋体')

    def h(text, level):
        return doc.add_heading(text, level=level)

    def para(text, bold=False, align=None):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        if align:
            p.alignment = align
        return p

    def bullets(items):
        for it in items:
            doc.add_paragraph(it, style='List Bullet')

    def figure(path, caption, width_cm=15.0):
        if path.exists():
            doc.add_picture(str(path), width=Cm(width_cm))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption)
        r.font.size = Pt(9)

    # ---------- 封面信息 ----------
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(f'{APP_NAME}\n软件说明书')
    r.bold = True
    r.font.size = Pt(22)
    for k, v in [('软件全称', f'{APP_NAME}'), ('软件简称', 'EchoHymn'),
                 ('版本号', f'V{version}'), ('著作权人', '赵金雪（Zhao Jinxue）'),
                 ('开发方式', '独立开发'), ('运行平台', 'Windows 10 / Windows 11（x64）')]:
        para(f'{k}：{v}')
    doc.add_page_break()

    # ---------- 1 概述 ----------
    h('1  软件概述', 1)
    h('1.1  简介', 2)
    para(f'{APP_NAME}是一款面向 Windows 桌面的赞美诗浏览与播放软件。软件内置诗歌库'
         '（含诗歌文本、曲谱数据与音频资源索引），提供歌词、曲谱、简谱、五线谱四种显示模式，'
         '并支持钢琴版/人声版音频播放、按节翻页与跟随播放自动翻页、歌单管理、全文检索、'
         '多套配色与全局字号缩放等功能。')
    h('1.2  主要功能', 2)
    bullets([
        '歌词显示：一页一节，底部翻页条，手动/自动翻页（自动跟随播放进度）；',
        '曲谱显示：简谱网格渲染（列 = 拍点），单声部主旋律与当前节歌词同列对齐，'
        '支持按节切换与跟随播放切节，内容不足一屏时整块垂直居中；',
        '简谱/五线谱显示：内置扫描图查看，支持缩放与滚动；',
        '音频播放：播放/暂停/上一首/下一首、进度拖动、音量调节、钢琴版与人声版切换；',
        '诗歌检索：按歌名与歌词检索，简体/繁体关键字双向匹配；',
        '歌单管理：内置分类歌单与个人歌单（新建/重命名/删除/增删成员）；',
        '个性化：五套配色主题切换、四级全局字号缩放、界面状态自动记忆。',
    ])
    h('1.3  运行环境', 2)
    bullets([
        '操作系统：Windows 10 / Windows 11（64 位）；',
        '硬件：常规 x64 桌面计算机，建议分辨率不低于 1280×800；',
        '网络：软件本体运行无需联网（数据随安装包本地分发）。',
    ])

    # ---------- 2 安装 ----------
    h('2  安装与卸载', 1)
    h('2.1  安装', 2)
    para('安装采用主体与素材分离的双产物分发：安装程序（EchoHymn_Setup_v%s.exe）与'
         '素材数据包（EchoHymn_Data_v%s.7z）须置于同一目录。运行安装程序后，'
         '安装向导依次进行环境检查、安装目录选择与安装确认，随后解出主程序；'
         '素材数据包由向导在同目录自动检测并解出至安装目录的数据子目录。' % (version, version))
    h('2.2  卸载', 2)
    para('通过"设置 → 应用"或开始菜单卸载条目卸载；卸载保留用户的个人歌单与界面状态数据。')

    # ---------- 3 界面总览 ----------
    h('3  界面总览', 1)
    para('主窗口自上而下分为：标题栏（软件名称与当前诗歌）、版本栏（音频版本切换与显示模式切换）、'
         '主内容区（左栏诗歌列表/歌单 + 右侧显示区）、播放条（进度与控制）与状态栏。'
         '版本栏右侧的四个按钮用于切换显示模式：歌词、曲谱、简谱、五线谱。')
    bullets([
        '左栏：诗歌列表、默认歌单、我的歌单三个栏目页签；',
        '显示区：按当前模式显示歌词 / 曲谱网格 / 简谱扫描图 / 五线谱扫描图；',
        '播放条：进度条、播放控制键、音量控件、人声版本列表入口。',
    ])

    # ---------- 4 显示模式 ----------
    h('4  显示模式操作说明', 1)
    h('4.1  歌词模式', 2)
    para('以整页文字显示当前节歌词（正歌与副歌分色）。底部翻页条以"◀ / 第 N / M 节 / ▶"切换节；'
         '右上角按钮在手动与自动翻页之间切换，自动模式下歌词随播放进度自动切节。')
    h('4.2  曲谱模式', 2)
    para('以简谱网格形式渲染曲谱：网格的每一个列对应一个拍点，音符、记号（减时线、高低八度点、'
         '延长记号）与小节线按列对齐，当前节歌词逐字落在对应列下方，实现谱与词的同列同步。'
         '曲谱按单声部（主旋律）显示；底部翻页条切换节，自动模式跟随播放进度切节；'
         '内容不足一屏时整块垂直居中，超出一屏时纵向滚动。')
    figure(img / 'fig2_score_single_voice.png',
           '图 1  曲谱模式：单声部简谱网格与当前节歌词同列对齐（第 9 首第 2 节）')
    para('网格按乐句块排布，每个乐句块以其自身列数铺满显示区宽度；'
         '小节线以竖线跨行绘制，保证乐句边界清晰。')
    figure(img / 'fig3_score_block_width.png',
           '图 2  曲谱模式：按乐句块宽度排布的版式（第 163 首）')
    h('4.3  简谱模式', 2)
    para('显示内置的印刷简谱扫描图。Ctrl+滚轮缩放（宽度自初始大小逐渐放大至铺满显示区），'
         '滚轮上下滚动；切换诗歌后自动复位为初始视图。')
    h('4.4  五线谱模式', 2)
    para('显示内置的五线谱扫描图，缩放与滚动操作同简谱模式。')

    # ---------- 5 播放 ----------
    h('5  音频播放操作说明', 1)
    bullets([
        '播放条中部为上一首 / 播放·暂停 / 下一首控制键，列表内循环；',
        '进度条可拖动跳转，两端显示已播放时间与总时长；',
        '左侧音量控件与系统音量双向同步（点击图标静音/取消）；',
        '版本栏可切换钢琴版与人声版；存在多个人声版本时，播放条提供版本列表入口。',
    ])

    # ---------- 6 检索与歌单 ----------
    h('6  诗歌检索与歌单管理', 1)
    para('顶栏搜索框支持按歌名与歌词检索，关键字简体/繁体双向匹配；命中歌词时定位到对应节。'
         '左栏"默认歌单"提供内置分类歌单；"我的歌单"支持新建、重命名、删除歌单及增删诗歌成员。')

    # ---------- 7 个性化 ----------
    h('7  个性化设置', 1)
    bullets([
        '配色：五套调色板主题一键切换，界面全部颜色随主题联动；',
        '字号：四级全局字号（默认/中号/大号/最大），整棵界面等比缩放；',
        '状态记忆：显示模式、翻页模式、配色、字号、歌单与播放位置等自动保存并于下次启动恢复。',
    ])

    # ---------- 8 数据 ----------
    h('8  数据与文件说明', 1)
    bullets([
        '安装目录 data/：诗歌与谱面数据库、音频与谱面素材资源；',
        '程序同级 state.json：界面状态与用户设置（便携存储）；',
        '程序同级 logs/：运行日志；',
        '曲谱网格数据为软件内置的结构化谱面数据（列 = 拍点），用于曲谱模式的渲染与谱词对齐。',
    ])
    para('说明：软件内展示的诗歌文本、曲谱图像与音频录音的著作权归相应权利人所有；'
         '本软件为上述内容的浏览与播放工具，其程序代码与文档的著作权归著作权人所有。')

    # ---------- 9 版本 ----------
    h('9  版本信息', 1)
    para(f'本说明书对应软件版本 V{version}。该版本要点：', bold=True)
    bullets([
        '曲谱模式采用简谱网格渲染，按单声部（主旋律）显示并与歌词同列对齐；',
        '曲谱支持一页一节与跟随播放自动切节；内容不足一屏时整块垂直居中；',
        '简谱网格按乐句块宽度排布，小节线跨行连续绘制；',
        '显示模式集为：歌词 / 曲谱 / 简谱 / 五线谱；',
        '谱面数据源为内置结构化网格数据（列 = 拍点）。',
    ])

    # ---------- 10 版权 ----------
    h('10  版权声明', 1)
    para('Copyright (c) 2026 赵金雪（Zhao Jinxue）。本软件程序代码与文档的著作权归著作权人所有；'
         '未经书面许可不得复制、修改、分发或用于商业用途（个人学习使用除外），详见随附许可协议。')

    return doc, OUTDIR / f'v{version}' / f'软件说明书_V{version}.docx'


if __name__ == '__main__':
    v = app_version()
    doc, out = build(v)
    doc.save(out)
    print('产物：', out)
